"""Excel-книга output/report.xlsx: P&L на формулах поверх проводок и маппинга.

Книга считает сама: лист «Ноги» разворачивает каждую проводку в дебетовую и кредитовую ногу, статья
подтягивается формулой из листа «Маппинг», P&L — SUMIFS по статье и месяцу. Финансист может
поменять правило в маппинге и увидеть, как изменится P&L. Тест пересчитывает книгу LibreOffice
и сверяет P&L с реестром до копейки.

Тексты из данных, начинающиеся с «=», пишутся строкой, а не формулой (защита от формул-инъекций).
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline.build_pl import MANAGEMENT_LINES

BOLD = Font(bold=True)
SUBTOTAL_FILL = PatternFill("solid", fgColor="EEF1F6")
MONEY = '#,##0.00;[Red]-#,##0.00'

HOW_TO_READ = [
    "Как читать книгу",
    "",
    "P&L — управленческий P&L по месяцам, рубли. Каждая ячейка статьи — формула SUMIFS по листу «Ноги»;",
    "промежуточные итоги — суммы статей выше. Числа совпадают с реестром чисел выпуска (registry.csv).",
    "Ноги — каждая проводка развёрнута на дебетовую (минус) и кредитовую (плюс) ногу; статья P&L",
    "подтягивается формулой из листа «Маппинг» по ключу «счёт|сторона|статья затрат».",
    "Проводки — исходная главная книга (journal.csv), как есть.",
    "Маппинг — правила разноски (mapping.csv); «тех» — закрытие счетов, в P&L не попадает.",
    "",
    "Данные вымышлены: демо-пример на синтетической главной книге.",
]


def _text(cell, value) -> None:
    """Записать строку как текст, даже если она начинается с «=»."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _month_date(month: str) -> date:
    return date(int(month[:4]), int(month[5:7]), 1)


def render_xlsx(out_dir: Path, data_dir: Path, months: list[str]) -> Path:
    journal = pd.read_csv(Path(data_dir) / "journal.csv", dtype=str, keep_default_na=False)
    mapping = pd.read_csv(Path(data_dir) / "mapping.csv", dtype=str, keep_default_na=False)

    wb = Workbook()
    how = wb.active
    how.title = "Как читать"
    for i, line in enumerate(HOW_TO_READ, 1):
        _text(how.cell(i, 1), line)
    how["A1"].font = Font(bold=True, size=14)
    how.column_dimensions["A"].width = 110

    pl = wb.create_sheet("P&L")
    legs = wb.create_sheet("Ноги")
    src = wb.create_sheet("Проводки")
    mp = wb.create_sheet("Маппинг")

    # Маппинг: правила + ключ
    mp.append(["Счёт", "Сторона", "Статья затрат", "Статья P&L", "Ключ"])
    for r, row in enumerate(mapping.itertuples(index=False), 2):
        for c, v in enumerate([row.account, row.side, row.cost_item, row.article], 1):
            _text(mp.cell(r, c), v)
        mp.cell(r, 5, f'=A{r}&"|"&B{r}&"|"&C{r}')
    last_mp = len(mapping) + 1

    # Проводки: как в journal.csv
    src.append(["№", "Дата", "Документ", "Дт", "Кт", "Сумма", "Статья затрат", "Содержание"])
    for r, row in enumerate(journal.itertuples(index=False), 2):
        src.cell(r, 1, int(row.entry_id))
        src.cell(r, 2, date.fromisoformat(row.date)).number_format = "yyyy-mm-dd"
        for c, v in [(3, row.doc), (4, row.debit), (5, row.credit), (7, row.cost_item), (8, row.description)]:
            _text(src.cell(r, c), v)
        src.cell(r, 6, float(row.amount)).number_format = MONEY

    # Ноги: две на проводку
    legs.append(["Месяц", "Счёт", "Сторона", "Статья затрат", "Статья P&L", "Сумма со знаком", "№ проводки"])
    r = 2
    for row in journal.itertuples(index=False):
        amount = float(row.amount)
        for account, side, signed in [(row.debit, "D", -amount), (row.credit, "C", amount)]:
            legs.cell(r, 1, _month_date(row.date[:7])).number_format = "yyyy-mm"
            _text(legs.cell(r, 2), account)
            _text(legs.cell(r, 3), side)
            _text(legs.cell(r, 4), row.cost_item)
            legs.cell(r, 5, f'=IFERROR(INDEX(Маппинг!$D$2:$D${last_mp},'
                            f'MATCH(B{r}&"|"&C{r}&"|"&D{r},Маппинг!$E$2:$E${last_mp},0)),"")')
            legs.cell(r, 6, signed).number_format = MONEY
            legs.cell(r, 7, int(row.entry_id))
            r += 1
    last_leg = r - 1

    # P&L
    pl["A1"] = "Управленческий P&L, руб. — демо-пример ООО «СтальКонструкт»"
    pl["A1"].font = Font(bold=True, size=13)
    pl.cell(3, 1, "Статья").font = BOLD
    for j, month in enumerate(months, 2):
        cell = pl.cell(3, j, _month_date(month))
        cell.number_format = "yyyy-mm"
        cell.font = BOLD
    article_rows: list[int] = []
    for i, (line, kind) in enumerate(MANAGEMENT_LINES, 4):
        _text(pl.cell(i, 1), line)
        for j in range(2, len(months) + 2):
            col = get_column_letter(j)
            if kind == "article":
                formula = (f"=SUMIFS(Ноги!$F$2:$F${last_leg},Ноги!$E$2:$E${last_leg},$A{i},"
                           f"Ноги!$A$2:$A${last_leg},{col}$3)")
            else:
                formula = "=" + "+".join(f"{col}{a}" for a in article_rows)
            cell = pl.cell(i, j, formula)
            cell.number_format = MONEY
        if kind == "article":
            article_rows.append(i)
        else:
            for j in range(1, len(months) + 2):
                pl.cell(i, j).font = BOLD
                pl.cell(i, j).fill = SUBTOTAL_FILL
    pl.column_dimensions["A"].width = 28
    for j in range(2, len(months) + 2):
        pl.column_dimensions[get_column_letter(j)].width = 17
    pl.freeze_panes = "B4"

    for ws, widths in [(legs, [10, 8, 8, 24, 24, 18, 11]), (src, [7, 11, 13, 7, 7, 16, 24, 48]),
                       (mp, [8, 8, 24, 24, 30])]:
        for k, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(k)].width = wdt
        for cell in ws[1]:
            cell.font = BOLD
        ws.freeze_panes = "A2"

    path = Path(out_dir) / "report.xlsx"
    wb.save(path)
    return path
