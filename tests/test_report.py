"""Отчёт: каждое число HTML есть в реестре; Excel-книга на формулах сходится с реестром до копейки."""
import shutil
import subprocess
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from pipeline.numbers import check_report_numbers, format_number, mln, pct, report_numbers
from pipeline.render_xlsx import render_xlsx
from pipeline.run import run
from tests.conftest import FIXTURE_JOURNAL, write_data

ROOT = Path(__file__).resolve().parents[1]


def find_soffice() -> str | None:
    for candidate in [shutil.which("soffice"), shutil.which("libreoffice"),
                      r"C:\Program Files\LibreOffice\program\soffice.exe"]:
        if candidate and Path(candidate).exists():
            return candidate
    return None


@pytest.fixture(scope="module")
def release(tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    assert run(ROOT / "data", out) == 0
    return out, pd.read_csv(out / "registry.csv", dtype={"period": str})


def test_format_russian():
    assert mln(4_547_712_345.67) == "4 547,7"
    assert mln(-1_450_000) == "−1,5"          # половина вверх, не банковское округление
    assert mln(370_400_000, signed=True) == "+370,4"
    assert pct(-0.25) == "−0,3"
    assert pct(0.0, signed=True) == "0,0"
    assert format_number(-10.5, "pct|neg|+") == "+10,5"
    assert format_number(-35_900_000, "mln|abs") == "35,9"


def test_parser_reads_numbers_and_skips_dates():
    text = ("<p>Выручка 4 547,7 млн ₽, −0,2 %, +8,9 %. Период 9М2026 к 9М2025, месяц 2026-09, "
            "сентябрь 2026, SHA-256 1f63e6254efdc64e…</p><style>.a{width:12px}</style>")
    values = [v for v, _ in report_numbers(text)]
    assert values == [Decimal("4547.7"), Decimal("-0.2"), Decimal("8.9")]


def broken_copy(out: Path, tmp_path: Path, old: str, new: str) -> Path:
    page = (out / "index.html").read_text(encoding="utf-8")
    assert old in page
    path = tmp_path / "index.html"
    path.write_text(page.replace(old, new, 1), encoding="utf-8")
    return path


def test_unbound_number_is_caught(release, tmp_path):
    out, registry = release
    path = broken_copy(out, tmp_path, "</main>", "<p>Выручка выросла на 12 345,6 млн ₽</p></main>")
    problems = check_report_numbers(path, registry)
    assert len(problems) == 1 and "12345.6" in problems[0]


def test_lost_sign_is_caught(release, tmp_path):
    out, registry = release
    row = registry[(registry["indicator"] == "Чистая прибыль") & (registry["measure"] == "изменение %")
                   & (registry["period"] == "9М2026 к 9М2025")].iloc[0]
    shown = format_number(row["value"], "pct|+")
    assert shown.startswith("−")
    span = f'data-r="{row["id"]}" data-f="pct|+">{shown}<'
    path = broken_copy(out, tmp_path, span, span.replace(shown, shown[1:]))
    problems = check_report_numbers(path, registry)
    assert len(problems) == 1 and row["id"] in problems[0]


def test_number_bound_to_wrong_row_is_caught(release, tmp_path):
    out, registry = release
    page = (out / "index.html").read_text(encoding="utf-8")
    rid = page.split('data-r="', 1)[1].split('"', 1)[0]
    other = "R0001" if rid != "R0001" else "R0002"
    path = broken_copy(out, tmp_path, f'data-r="{rid}"', f'data-r="{other}"')
    assert any(other in p for p in check_report_numbers(path, registry))


def test_release_report_matches_registry(release):
    out, registry = release
    assert check_report_numbers(out / "index.html", registry) == []
    page = (out / "index.html").read_text(encoding="utf-8")
    assert "СтальКонструкт-Демо" in page and "Отклонения сверх роста выручки" in page
    assert len(report_numbers(page)) > 100


def test_xlsx_structure_and_text_injection(tmp_path):
    journal = [r for r in FIXTURE_JOURNAL]
    data = write_data(tmp_path / "data", journal)
    j = pd.read_csv(data / "journal.csv", dtype=str, keep_default_na=False)
    j.loc[0, "description"] = "=HYPERLINK(\"http://example.com\")"
    j.to_csv(data / "journal.csv", index=False, lineterminator="\n")
    path = render_xlsx(tmp_path, data, ["2025-01"])
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Как читать", "P&L", "Ноги", "Проводки", "Маппинг"]
    cell = wb["Проводки"]["H2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    assert wb["P&L"]["B4"].value.startswith("=SUMIFS(")


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice не установлен")
def test_xlsx_recalculates_to_registry(release, tmp_path):
    out, registry = release
    subprocess.run([find_soffice(), "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_path),
                    str(out / "report.xlsx")], check=True, capture_output=True, timeout=300)
    wb = openpyxl.load_workbook(tmp_path / "report.xlsx", data_only=True)
    ws = wb["P&L"]
    months = [c.value.strftime("%Y-%m") for c in ws[3][1:] if c.value is not None]
    monthly = registry[(registry["measure"] == "значение") & registry["period"].str.fullmatch(r"\d{4}-\d{2}")]
    expected = monthly.set_index(["indicator", "period"])["value"]
    errors, compared = [], 0
    for row in ws.iter_rows(min_row=4):
        line = row[0].value
        if not line:
            continue
        for month, cell in zip(months, row[1:]):
            assert not (isinstance(cell.value, str) and cell.value.startswith("#")), (line, month, cell.value)
            if abs(cell.value - expected[(line, month)]) > 0.005:
                errors.append((line, month, cell.value, expected[(line, month)]))
            compared += 1
    assert compared == len(expected)
    assert errors == []
